#!/usr/bin/env python3
"""Build the maintained Smart Citizen manuscript with a small explicit Markdown subset.

Run with the primary artifact runtime Python. Requires python-docx; SVG embedding
uses the primary runtime Node and sharp identified by CODEX_PRIMARY_RUNTIME_NODE
and CODEX_PRIMARY_RUNTIME_NODE_MODULES. No network or repository npm installation
is needed. The artifact-operation marker belongs to the authoring workflow and
must be run by that workflow before the first DOCX build, not by this script.

Supported source: ATX headings, prose, inline emphasis/code/links, simple lists,
pipe tables, fenced listings, standalone local images, and Figure/Table/Listing
captions. The References section uses hanging indents. HTML is intentionally not
used as a formatting language. Author-date citations are retained as authored.

Examples:
  "$CODEX_PRIMARY_RUNTIME_PYTHON" scripts/build-smart-citizen-paper.py --figures-only
  "$CODEX_PRIMARY_RUNTIME_PYTHON" scripts/build-smart-citizen-paper.py
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE


ROOT = Path(__file__).resolve().parents[1]
STUDY = ROOT / "Docs/Research/Studies/Smart_Citizen"
DEFAULT_SOURCE = STUDY / "MANUSCRIPT.md"
DEFAULT_OUTPUT = STUDY / "SMART_CITIZEN_EXPERIMENT_PAPER.docx"
BODY_WIDTH = 6.8
BLACK = RGBColor(0, 0, 0)


def svg_start(title: str, height: int) -> list[str]:
    return [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="{height}" viewBox="0 0 1200 {height}" role="img" aria-labelledby="title desc">',
        f'<title id="title">{escape(title)}</title>',
        '<desc id="desc">Proposed study design. Boxes and arrows describe intended roles and relationships, not an implemented service or verified physical installation.</desc>',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="4" orient="auto"><path d="M0 0 L8 4 L0 8 Z" fill="#333"/></marker></defs>',
        '<style>text{font-family:Arial,Liberation Sans,sans-serif;fill:#111}.title{font-size:32px;font-weight:700}.head{font-size:25px;font-weight:700}.body{font-size:24px}.label{font-size:23px}.note{font-size:23px;font-style:italic}.box{fill:#f5f5f5;stroke:#555;stroke-width:2}.edge{fill:none;stroke:#333;stroke-width:2.5;marker-end:url(#arrow)}.uncertain{stroke-dasharray:9 7}</style>',
        '<rect x="0" y="0" width="1200" height="100%" fill="#fff"/>',
        f'<text class="title" x="600" y="44" text-anchor="middle">{escape(title)}</text>',
    ]


def svg_box(parts: list[str], x: int, y: int, w: int, h: int, title: str, lines: list[str], dashed=False):
    klass = "box uncertain" if dashed else "box"
    parts.append(f'<rect class="{klass}" x="{x}" y="{y}" width="{w}" height="{h}" rx="5"/>')
    center = x + w / 2
    parts.append(f'<text class="head" x="{center:g}" y="{y+35}" text-anchor="middle">{escape(title)}</text>')
    for j, line in enumerate(lines):
        parts.append(f'<text class="body" x="{center:g}" y="{y+72+j*30}" text-anchor="middle">{escape(line)}</text>')


def svg_edge(parts: list[str], d: str, label: str = "", at: tuple[int, int] | None = None, dashed=False):
    parts.append(f'<path class="edge{" uncertain" if dashed else ""}" d="{d}"/>')
    if label and at:
        x, y = at
        parts.append(f'<text class="label" x="{x}" y="{y}" text-anchor="middle">{escape(label)}</text>')


def diagram_architecture() -> str:
    p = svg_start("Proposed experiment architecture", 1020)
    svg_box(p, 60, 90, 320, 135, "Smart Citizen API", ["Resource metadata", "and observations"])
    svg_box(p, 440, 90, 320, 135, "Published definitions", ["API and hardware", "documentation"])
    svg_box(p, 820, 90, 320, 135, "Target definitions", ["OGC representations", "and MPSS bindings"])
    svg_box(p, 60, 310, 320, 135, "Captured evidence", ["Reviewed source data", "and acquisition record"])
    svg_box(p, 440, 310, 320, 135, "Adapter and mapping", ["Explicit transformations", "and unresolved fields"])
    svg_box(p, 820, 310, 320, 135, "Mapping record", ["Source to target trace", "and versioned decisions"])
    svg_box(p, 440, 525, 320, 120, "CSAPI test service", ["Descriptions and data"])
    svg_box(p, 60, 740, 320, 120, "Vendor interface", ["Source presentation"])
    svg_box(p, 440, 740, 320, 120, "Generic reader", ["Competent standard view"])
    svg_box(p, 820, 740, 320, 120, "Profile aware view", ["Bounded interpretation"])
    svg_box(p, 310, 915, 580, 70, "Paired evidence and evaluation", [])
    svg_edge(p, "M220 225 V310")
    svg_edge(p, "M600 225 V310")
    svg_edge(p, "M980 225 V265 H730 V310")
    svg_edge(p, "M380 377 H440")
    svg_edge(p, "M760 377 H820")
    svg_edge(p, "M600 445 V525")
    svg_edge(p, "M600 645 V740")
    svg_edge(p, "M760 585 H980 V740")
    svg_edge(p, "M60 157 H25 V800 H60", dashed=True)
    svg_edge(p, "M220 860 V950 H310")
    svg_edge(p, "M600 860 V915")
    svg_edge(p, "M980 860 V950 H890")
    p.append('<text class="note" x="220" y="682" text-anchor="middle">Reference view of source data</text>')
    p.append('</svg>')
    return "\n".join(p) + "\n"


def diagram_model() -> str:
    p = svg_start("Proposed distinctions among represented subjects", 990)
    svg_box(p, 60, 100, 340, 135, "Model description", ["Reusable capabilities", "and technical description"])
    svg_box(p, 800, 100, 340, 135, "Source registration", ["Platform record identifiers", "and reported metadata"])
    svg_box(p, 430, 320, 340, 135, "Represented system", ["Physical identity policy", "must be explicit"], dashed=True)
    svg_box(p, 800, 320, 340, 135, "Adapter provenance", ["Representation source", "and mapping revision"])
    svg_box(p, 60, 580, 340, 135, "Data stream", ["Output and result schema", "with qualified context"])
    svg_box(p, 430, 580, 340, 135, "Observation", ["Value and time", "with source provenance"])
    svg_box(p, 800, 580, 340, 135, "Observed subject", ["Environment or equipment", "as supported by evidence"], dashed=True)
    svg_edge(p, "M230 235 V387 H430", "model applies", (124, 292))
    svg_edge(p, "M970 235 V275 H600 V320", "source representation", (775, 266), dashed=True)
    svg_edge(p, "M800 387 H770")
    svg_edge(p, "M430 420 H230 V580", "produces", (310, 523))
    svg_edge(p, "M400 647 H430")
    svg_edge(p, "M770 647 H800")
    p.append('<text class="label" x="600" y="552" text-anchor="middle">Stream supplies records</text>')
    p.append('<text class="label" x="970" y="552" text-anchor="middle">Record is about a subject</text>')
    p.append('<path d="M80 790 H1120" stroke="#aaa" stroke-width="1.5"/>')
    p.append('<text class="head" x="600" y="837" text-anchor="middle">Power information retains its limits</text>')
    p.append('<text class="body" x="600" y="880" text-anchor="middle">Battery percentage does not identify a physical battery</text>')
    p.append('<text class="body" x="600" y="913" text-anchor="middle">or establish a directed supply relationship.</text>')
    p.append('<text class="note" x="600" y="965" text-anchor="middle">Dashed elements require an explicit identification or subject decision.</text>')
    p.append('</svg>')
    return "\n".join(p) + "\n"


def write_figures(directory: Path):
    directory.mkdir(parents=True, exist_ok=True)
    definitions = {
        "figure_01_architecture.svg": diagram_architecture(),
        "figure_02_model.svg": diagram_model(),
    }
    for name, content in definitions.items():
        (directory / name).write_text(content, encoding="utf-8")
    return [directory / name for name in definitions]


def set_font(style, name: str, size: float, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ["ascii", "hAnsi", "cs", "eastAsia"]:
        fonts.set(qn("w:" + key), name)
    for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn("w:" + key), None)


def configure_document(doc, version: str, revision_date: date):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = Inches(0.77)
    sec.bottom_margin = Inches(0.77)
    sec.header_distance = sec.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    set_font(normal, "Times New Roman", 11.5)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 20), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11.5), ("Heading 4", 11.5)]:
        s = doc.styles[name]
        set_font(s, "Times New Roman", size, bold=True)
        s.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.keep_together = True
        ppr = s.element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    styles = {
        "Paper Caption": ("Times New Roman", 10, False, False),
        "Paper Code": ("DejaVu Sans Mono", 9, False, False),
        "Paper Reference": ("Times New Roman", 10.5, False, False),
        "Paper List": ("Times New Roman", 11.5, False, False),
    }
    for name, (font, size, bold, italic) in styles.items():
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        set_font(style, font, size, bold, italic)
    code = doc.styles["Paper Code"].paragraph_format
    code.space_before = code.space_after = Pt(0)
    code.line_spacing = 1.0
    code.left_indent = Inches(0.07)
    code.right_indent = Inches(0.07)
    ref = doc.styles["Paper Reference"].paragraph_format
    ref.left_indent, ref.first_line_indent = Inches(0.28), Inches(-0.28)
    ref.space_after = Pt(6)
    ref.line_spacing = 1.05
    ref.keep_together = True
    cap = doc.styles["Paper Caption"].paragraph_format
    cap.space_before, cap.space_after = Pt(3), Pt(9)
    cap.line_spacing = 1.05
    cap.keep_together = True
    lst = doc.styles["Paper List"].paragraph_format
    lst.left_indent, lst.first_line_indent = Inches(0.18), Inches(-0.18)
    lst.space_after = Pt(4)
    header = sec.header.paragraphs[0]
    header.text = "Smart Citizen integration study"
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        run.font.name, run.font.size, run.font.color.rgb = "Times New Roman", Pt(9), BLACK
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"Research draft {version}   {revision_date.day} {revision_date.strftime('%B %Y')} EDT   ")
    r.font.name, r.font.size, r.font.color.rgb = "Times New Roman", Pt(9), BLACK
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    props = doc.core_properties
    props.author = "MERLIN Project"
    props.subject = "Proposed Smart Citizen integration experiment"
    props.version = version
    stamp = datetime(revision_date.year, revision_date.month, revision_date.day, tzinfo=timezone.utc)
    props.created = props.modified = stamp


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^\n)]+\)|<https?://[^>]+>)")


def add_hyperlink(paragraph, text: str, url: str, italic=False, bold=False):
    link = OxmlElement("w:hyperlink")
    if url.startswith("#"):
        link.set(qn("w:anchor"), url[1:])
    else:
        rid = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    if italic:
        props.append(OxmlElement("w:i"))
    if bold:
        props.append(OxmlElement("w:b"))
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text: str, source_dir: Path, bold=None, italic=None):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(9.5)
            run.bold, run.italic = bold, italic
        elif piece.startswith("**") and piece.endswith("**"):
            add_inline(paragraph, piece[2:-2], source_dir, bold=True, italic=italic)
        elif piece.startswith("*") and piece.endswith("*"):
            add_inline(paragraph, piece[1:-1], source_dir, bold=bold, italic=True)
        elif piece.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^\n)]+)\)", piece)
            label, url = match.groups()
            if not re.match(r"[a-zA-Z][a-zA-Z0-9+.-]*:", url) and not url.startswith("#"):
                target = (source_dir / url).resolve()
                if target.is_relative_to(ROOT):
                    url = "https://github.com/OS4CSAPI/glaux-client/blob/work/details-view-20261023/" + target.relative_to(ROOT).as_posix()
            label_italic = italic or (label.startswith("*") and label.endswith("*"))
            label_bold = bold or (label.startswith("**") and label.endswith("**"))
            add_hyperlink(paragraph, label.strip("*"), url, label_italic, label_bold)
        elif piece.startswith("<http") and piece.endswith(">"):
            add_hyperlink(paragraph, piece[1:-1], piece[1:-1], italic, bold)
        else:
            run = paragraph.add_run(piece.replace("\\|", "|"))
            run.bold, run.italic = bold, italic


def table_cells(line: str) -> list[str]:
    return [x.strip() for x in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def add_table(doc, lines: list[str], source_dir: Path):
    rows = [table_cells(x) for x in lines]
    headers, body = rows[0], rows[2:]
    cols = len(headers)
    if any(len(row) != cols for row in body):
        raise ValueError("Inconsistent manuscript table column count")
    widths = {2: [2.0, 4.8], 3: [1.55, 2.3, 2.95], 4: [1.2, 1.3, 1.15, 3.15], 5: [0.6, 1.25, 1.5, 1.7, 1.75]}.get(cols, [BODY_WIDTH / cols] * cols)
    if cols == 3 and headers[0].strip("*").lower() == "figure":
        widths = [0.6, 3.15, 3.05]
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for row_index, values in enumerate([headers, *body]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, (cell, content) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                border = OxmlElement("w:" + edge)
                for k, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    border.set(qn("w:" + k), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                entry = OxmlElement("w:" + edge)
                entry.set(qn("w:w"), "90")
                entry.set(qn("w:type"), "dxa")
                margins.append(entry)
            props.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8E8E8" if row_index == 0 else ("F7F7F7" if row_index % 2 == 0 else "FFFFFF"))
            props.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.03
            add_inline(p, content, source_dir)
            for run in p.runs:
                run.font.size = Pt(10)
                if row_index == 0:
                    run.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(3)


def rasterize_svg(source: Path, target: Path):
    node = os.environ.get("CODEX_PRIMARY_RUNTIME_NODE")
    modules = os.environ.get("CODEX_PRIMARY_RUNTIME_NODE_MODULES")
    if not node or not modules:
        raise RuntimeError("SVG embedding requires the primary runtime Node and Node modules environment variables")
    script = "const sharp=require(process.env.CODEX_PRIMARY_RUNTIME_NODE_MODULES+'/sharp');sharp(process.argv[1],{density:300}).resize({width:2200}).png().toFile(process.argv[2]).catch(e=>{console.error(e);process.exit(1)});"
    subprocess.run([node, "-e", script, str(source), str(target)], check=True)


def add_image(doc, alt: str, image_path: str, source_dir: Path, temporary: Path):
    path = (source_dir / image_path).resolve()
    if not path.is_file() or not path.is_relative_to(ROOT):
        raise ValueError(f"Missing or non-repository manuscript image: {image_path}")
    embed = path
    if path.suffix.lower() == ".svg":
        embed = temporary / (path.stem + ".png")
        rasterize_svg(path, embed)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    shape = p.add_run().add_picture(str(embed), width=Inches(BODY_WIDTH))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)


def normalized_docx_save(doc, output: Path):
    raw = io.BytesIO()
    doc.save(raw)
    source = zipfile.ZipFile(raw)
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
        for old in source.infolist():
            info = zipfile.ZipInfo(old.filename, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = old.external_attr
            destination.writestr(info, source.read(old.filename))


def build(source: Path, output: Path, version: str, revision_date: date):
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_document(doc, version, revision_date)
    references = False
    title_seen = False
    i = 0
    with tempfile.TemporaryDirectory(prefix="glaux-paper-") as tmp:
        temporary = Path(tmp)
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            if line.startswith("<!--"):
                while i < len(lines) and "-->" not in lines[i]:
                    i += 1
                i += 1
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                level, content = len(heading[1]), heading[2]
                if level == 1 and not title_seen:
                    p = doc.add_paragraph(style="Title")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    doc.core_properties.title = content
                    title_seen = True
                else:
                    p = doc.add_paragraph(style=f"Heading {max(1,min(level-1,4))}")
                add_inline(p, content, source.parent)
                references = bool(re.search(r"\bReferences$", content))
                i += 1
                continue
            if line.startswith("```"):
                listing = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    listing.append(lines[i])
                    i += 1
                if i >= len(lines):
                    raise ValueError("Unclosed code fence")
                for n, code_line in enumerate(listing):
                    p = doc.add_paragraph(style="Paper Code")
                    p.add_run(code_line)
                    p.paragraph_format.keep_with_next = n < len(listing) - 1
                if listing:
                    p.paragraph_format.space_after = Pt(8)
                i += 1
                continue
            image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if image:
                add_image(doc, image[1], image[2], source.parent, temporary)
                i += 1
                continue
            if i + 1 < len(lines) and "|" in line and re.fullmatch(r"[\s|:\-]+", lines[i+1]) and "---" in lines[i+1]:
                table_lines = [lines[i], lines[i+1]]
                i += 2
                while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                    table_lines.append(lines[i])
                    i += 1
                add_table(doc, table_lines, source.parent)
                continue
            if line in ["---", "***"]:
                i += 1
                continue
            bullet = re.match(r"^(?:([-*+])|([0-9]+)[.)])\s+(.+)$", line)
            if bullet:
                p = doc.add_paragraph(style="Paper List")
                prefix = "• " if bullet[1] else bullet[2] + ". "
                p.add_run(prefix)
                text_line = bullet[3]
                i += 1
                while i < len(lines) and lines[i].startswith("  ") and lines[i].strip():
                    text_line += " " + lines[i].strip()
                    i += 1
                add_inline(p, text_line, source.parent)
                continue
            # A prose paragraph comprises consecutive source lines separated by
            # a blank line or a supported block marker.
            prose = [line]
            i += 1
            while i < len(lines) and lines[i].strip():
                candidate = lines[i].strip()
                if re.match(r"^(#|```|!\[|<!--|[-*+]\s|[0-9]+[.)]\s|>)", candidate):
                    break
                if i + 1 < len(lines) and "|" in candidate and re.fullmatch(r"[\s|:\-]+", lines[i+1]):
                    break
                prose.append(candidate)
                i += 1
            content = " ".join(prose)
            if content.startswith("> "):
                content = content[2:]
            caption = re.match(r"^(?:\*\*)?(Figure|Table|Listing)\s+[0-9]+[.:\s]", content)
            style = "Paper Reference" if references else ("Paper Caption" if caption else "Normal")
            p = doc.add_paragraph(style=style)
            if caption:
                p.paragraph_format.keep_with_next = caption[1] in ["Table", "Listing"]
            add_inline(p, content, source.parent)
        input_paths = {source, Path(__file__).resolve()}
        for image_path in re.findall(r"^!\[[^\]]*\]\(([^)]+)\)\s*$", text, flags=re.M):
            input_paths.add((source.parent / image_path).resolve())
        input_digests = {
            path.relative_to(ROOT).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted(input_paths)
        }
        canonical_inputs = json.dumps(input_digests, sort_keys=True, separators=(",", ":")).encode("utf-8")
        doc.core_properties.comments = json.dumps({
            "format": "glaux-paper-inputs-v1",
            "sha256": hashlib.sha256(canonical_inputs).hexdigest(),
        }, sort_keys=True, separators=(",", ":"))
        normalized_docx_save(doc, output)
    return hashlib.sha256(output.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--version", default="0.1")
    parser.add_argument("--date", default="2026-10-24", help="Editorial revision date in America/New_York, YYYY-MM-DD")
    parser.add_argument("--figures-only", action="store_true", help="Regenerate the two authored SVG diagrams without building a DOCX")
    parser.add_argument("--refresh-figures", action="store_true", help="Regenerate SVG diagrams before building the document")
    args = parser.parse_args()
    if args.figures_only or args.refresh_figures:
        paths = write_figures(args.source.parent / "figures")
        for path in paths:
            print(f"Wrote {path.relative_to(ROOT)}")
    if not args.figures_only:
        digest = build(args.source.resolve(), args.output.resolve(), args.version, date.fromisoformat(args.date))
        print(f"Wrote {args.output}\nSHA-256 {digest}")


if __name__ == "__main__":
    main()
